# %%
# src/talaria/hermes_service/rag_processor/parser.py

import logging
from pathlib import Path
from typing import Optional, List
import re
from docx import Document
from pptx import Presentation
import fitz

# %%
log = logging.getLogger(__name__)


def _parse_pdf(file_path: Path) -> Optional[str]:
    """ pdf 파일에서 텍스트를 추출합니다. """
    try:
        text_parts: List[str] = []
        doc = fitz.open(file_path) # fitz.open() 사용
        
        # 모든 페이지 순회
        for page in doc:
            # "text" 옵션은 pypdf2보다 훨씬 더 나은 텍스트 순서를 반환합니다.
            text = page.get_text("text") 
            if text:
                text_parts.append(text)
        
        doc.close()
        
        raw_text = "\n\n".join(text_parts) # 페이지별로 두 줄 띄어쓰기

        # 1. 3개 이상의 연속된 줄바꿈을 "진짜 문단 구분"인 \n\n 으로 통일
        cleaned_text = re.sub(r'\n{3,}', '\n\n', raw_text)
        
        # 2. (핵심) 문단 구분이 아닌 "단일" 줄바꿈(\n)만 공백(' ')으로 변경
        # cleaned_text = re.sub(r'(?<!\n)\n(?!\n)', ' ', cleaned_text)
        
        # # 3. '공백/탭' 정리
        cleaned_text = re.sub(r'[ \t]+', ' ', cleaned_text).strip()

        
        
        return cleaned_text
        
    except Exception as e:
        log.error(f".pdf(fitz) 파일 파싱 실패 ({file_path.name}): {e}")
        return None


def _parse_docx(file_path: Path) -> Optional[str]:
    """ .docx (Word) 파일에서 '문단'과 '표' 텍스트를 모두 추출합니다. """
    try:
        text_parts: List[str] = []
        doc = Document(file_path)
        
        # --- 1. (기존) 문단(Paragraphs) 텍스트 추출 ---
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if para_text:
                para_text = re.sub(r'[ \t]+', ' ', para_text)
                text_parts.append(para_text)
        
        # --- 2. 표(Tables) 텍스트 추출 ---
        for table in doc.tables:
            table_text_parts: List[str] = []
            for row in table.rows:
                row_text_parts: List[str] = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        # 셀 안의 텍스트도 정리 (줄바꿈을 공백으로)
                        cell_text = re.sub(r'[\n\t]+', ' ', cell_text)
                        cell_text = re.sub(r' +', ' ', cell_text)
                        row_text_parts.append(cell_text)
                
                # 각 셀의 텍스트를 " | " (파이프)로 구분하여 한 줄로 만듭니다.
                if row_text_parts:
                    table_text_parts.append(" | ".join(row_text_parts))
            
            # 표(table) 하나가 끝나면, 그 내용을 하나의 블록으로 추가합니다.
            if table_text_parts:
                text_parts.append("\n".join(table_text_parts))

        # 3. (기존) 문단과 표 텍스트를 모두 \n\n (문단 구분)으로 합칩니다.
        return "\n\n".join(text_parts)
        
    except Exception as e:
        log.error(f".docx 파일 파싱 실패 ({file_path.name}): {e}")
        return None


def _parse_txt(file_path: Path) -> Optional[str]:
    """ .txt 파일에서 텍스트를 추출하고 '문단'을 보존합니다. """
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            raw_text = f.read()
    except UnicodeDecodeError:
        log.warning(f"UTF-8 디코딩 실패. '{file_path.name}' 파일을 latin-1로 재시도합니다.")
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                raw_text = f.read()
        except Exception as e:
            log.error(f".txt 파일 읽기 실패 ({file_path.name}): {e}")
            return None
    except Exception as e:
        log.error(f".txt 파일 읽기 실패 ({file_path.name}): {e}")
        return None

    # .txt의 문단 구분을 존중합니다.
    cleaned_text = re.sub(r'\n{3,}', '\n\n', raw_text)
    cleaned_text = re.sub(r'[ \t]+', ' ', cleaned_text).strip()
    return cleaned_text


# --- 메인 파서 함수 (이 함수는 수정할 필요가 없습니다) ---
def parse_document(file_path: str) -> Optional[str]:
    # ... (이전과 동일) ...
    try:
        path_obj = Path(file_path)
        file_ext = path_obj.suffix.lower()

        log.info(f"문서 파싱 시작: {path_obj.name}")

        if not path_obj.exists():
            log.error(f"파일을 찾을 수 없습니다: {file_path}")
            return None

        # 파일 확장자에 따라 적절한 파서 호출
        if file_ext == '.pdf':
            return _parse_pdf(path_obj)
        elif file_ext == '.docx':
            return _parse_docx(path_obj)
        # elif file_ext == '.pptx':
        #     return _parse_pptx(path_obj) # (pptx도 위 코드의 _parse_pptx로 교체해주세요)
        elif file_ext == '.txt':
            return _parse_txt(path_obj)
        else:
            log.warning(f"지원하지 않는 파일 형식입니다: {file_ext}")
            return None

    except Exception as e:
        log.error(f"'{file_path}' 문서 처리 중 예기치 않은 오류 발생: {e}", exc_info=True)
        return None